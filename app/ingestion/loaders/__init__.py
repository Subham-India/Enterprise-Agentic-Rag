from pathlib import Path

import pdfplumber
from bs4 import BeautifulSoup
from docx import Document
from pptx import Presentation
from pypdf import PdfReader


def load_pdf(path: Path) -> str:
    # Fast path: pypdf
    try:
        text = "\n\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
        if len(text.strip()) > 100:
            return text
    except Exception:
        pass
    # Fallback: pdfplumber
    with pdfplumber.open(str(path)) as pdf:
        return "\n\n".join((p.extract_text() or "") for p in pdf.pages)


def load_html(path: Path) -> str:
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    for tag in soup(["script", "style", "nav", "footer"]):
        tag.decompose()
    return soup.get_text("\n")


def load_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def load_docx(path: Path) -> str:
    doc = Document(str(path))
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())


def load_pptx(path: Path) -> str:
    parts = []
    for slide in Presentation(str(path)).slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
    return "\n\n".join(parts)


LOADERS = {
    ".pdf": load_pdf,
    ".html": load_html,
    ".txt": load_txt,
    ".docx": load_docx,
    ".pptx": load_pptx,
}


def load_file(path: Path) -> str | None:
    """Return the file's text, or None if the file type isn't supported."""
    loader = LOADERS.get(path.suffix.lower())
    return loader(path) if loader else None